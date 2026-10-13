import io
import pytest
from fastapi.testclient import TestClient
from docx import Document
from sqlalchemy import text

from backend.main import app
from backend.app.database import init_db, SessionLocal, engine


@pytest.fixture(scope="session", autouse=True)
def setup_db():
    init_db()


@pytest.fixture
def client():
    with TestClient(app) as c:
        res = c.post("/api/auth/login", json={"email": "alex@example.com", "password": "student123"})
        if res.status_code == 200:
            token = res.json()["access_token"]
            c.headers["Authorization"] = f"Bearer {token}"
        yield c


def create_sample_docx_bytes() -> bytes:
    """Helper to generate a realistic in-memory DOCX file."""
    doc = Document()
    doc.add_heading("Cloud Computing & Distributed Systems", level=1)
    doc.add_paragraph(
        "Distributed systems consist of autonomous computing entities that communicate over a network. "
        "Virtualization provides abstraction of physical computing resources, including CPU, storage, and networking. "
        "Microservices architecture structures an application as a collection of loosely coupled services. "
        "Load balancing distributes incoming network traffic across multiple servers to prevent overload. "
        "Fault tolerance ensures that a system continues operating properly in the event of component failure. "
        "Containerization with Docker encapsulates software and dependencies into portable execution units."
    )
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_complete_12_step_user_workflow(client):
    """
    SECTION 2 VERIFICATION: Complete End-to-End User Workflow Test
    1. Open EduMind AI (verify health & root info).
    2. Upload a DOCX/PDF learning material.
    3. Verify text is extracted and saved in SQLite.
    4. Generate a quiz from the uploaded material.
    5. Open the generated quiz and start an attempt (verify answers protected).
    6. Answer questions and submit the attempt.
    7. Verify backend calculates and saves result correctly.
    8. Open result and review answers and explanations.
    9. Open attempt history and verify saved attempts appear.
    10. Open competency analysis and verify it uses real quiz results.
    11. Open personalized recommendations and verify based on identified gaps.
    12. Verify persistence across database reconnections.
    """

    # -------------------------------------------------------------------------
    # STEP 1: Open EduMind AI & Health Check
    # -------------------------------------------------------------------------
    root_res = client.get("/")
    assert root_res.status_code == 200
    assert root_res.json()["project"] == "EduMind AI"

    health_res = client.get("/api/health")
    assert health_res.status_code == 200
    health_data = health_res.json()
    assert health_data["status"] == "healthy"
    assert health_data["database"]["status"] == "connected"

    # -------------------------------------------------------------------------
    # STEP 2: Upload Learning Material (DOCX)
    # -------------------------------------------------------------------------
    docx_bytes = create_sample_docx_bytes()
    upload_res = client.post(
        "/api/materials/upload",
        files={"file": ("distributed_cloud.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={
            "title": "Distributed Cloud Systems",
            "subject_code": "CS401",
            "description": "Lecture notes on virtualization, microservices, and load balancing."
        }
    )
    assert upload_res.status_code == 201, f"Upload failed: {upload_res.text}"
    uploaded_material = upload_res.json()
    material_id = uploaded_material["id"]
    assert material_id > 0
    assert uploaded_material["file_type"] == "docx"
    assert uploaded_material["status"] == "ready"

    # -------------------------------------------------------------------------
    # STEP 3: Verify Text Extracted and Saved
    # -------------------------------------------------------------------------
    detail_res = client.get(f"/api/materials/{material_id}")
    assert detail_res.status_code == 200
    assert detail_res.json()["extracted_char_count"] > 100

    content_res = client.get(f"/api/materials/{material_id}/content")
    assert content_res.status_code == 200
    extracted_text = content_res.json()["extracted_text"]
    assert "Virtualization provides abstraction" in extracted_text
    assert "Microservices architecture" in extracted_text

    # -------------------------------------------------------------------------
    # STEP 4: Generate Quiz from Uploaded Material
    # -------------------------------------------------------------------------
    gen_res = client.post(
        "/api/quizzes/generate",
        json={
            "material_id": material_id,
            "title": "Cloud Virtualization Assessment",
            "topic": "Virtualization & Containers",
            "num_questions": 5,
            "difficulty": "Medium"
        }
    )
    assert gen_res.status_code == 201, f"Quiz generation failed: {gen_res.text}"
    quiz_data = gen_res.json()
    quiz_id = quiz_data["id"]
    assert quiz_id > 0
    assert quiz_data["total_questions"] == 5

    # -------------------------------------------------------------------------
    # STEP 5: Open Generated Quiz & Start Attempt (Verify Answers Protected)
    # -------------------------------------------------------------------------
    quiz_get_res = client.get(f"/api/quizzes/{quiz_id}")
    assert quiz_get_res.status_code == 200

    start_res = client.post(f"/api/quizzes/{quiz_id}/attempts")
    assert start_res.status_code == 201
    attempt_init = start_res.json()
    attempt_id = attempt_init["attempt_id"]
    questions = attempt_init["questions"]
    assert len(questions) == 5

    # Security check: correct_index and explanation MUST NOT be returned
    for q in questions:
        assert "correct_index" not in q, "Security leak: correct_index present in start attempt!"
        assert "explanation" not in q, "Security leak: explanation present in start attempt!"
        assert len(q["options"]) == 4

    # -------------------------------------------------------------------------
    # STEP 6: Answer Questions and Submit Attempt
    # -------------------------------------------------------------------------
    # We deliberately submit answers to achieve an intentional score:
    # 1 correct, 4 incorrect to simulate a gap in "Virtualization & Containers"
    # Inspect server-side quiz question correct indices directly to control test responses:
    with SessionLocal() as db:
        from backend.app.models.quiz import QuizQuestion
        db_questions = db.query(QuizQuestion).filter(QuizQuestion.quiz_id == quiz_id).order_by(QuizQuestion.order_num).all()
        q_map = {q.id: q.correct_index for q in db_questions}

    responses_payload = []
    # Q1: answer correctly
    responses_payload.append({
        "question_id": questions[0]["id"],
        "selected_index": q_map[questions[0]["id"]]
    })
    # Q2 to Q5: answer incorrectly
    for q in questions[1:]:
        wrong_index = (q_map[q["id"]] + 1) % 4
        responses_payload.append({
            "question_id": q["id"],
            "selected_index": wrong_index
        })

    submit_res = client.post(
        f"/api/attempts/{attempt_id}/submit",
        json={
            "responses": responses_payload,
            "time_taken_seconds": 125
        }
    )
    assert submit_res.status_code == 200, f"Submit failed: {submit_res.text}"
    result_data = submit_res.json()

    # -------------------------------------------------------------------------
    # STEP 7: Verify Backend Calculates and Saves Result Correctly
    # -------------------------------------------------------------------------
    assert result_data["status"] == "completed"
    assert result_data["total_questions"] == 5
    assert result_data["attempted_count"] == 5
    assert result_data["correct_count"] == 1
    assert result_data["incorrect_count"] == 4
    assert result_data["unanswered_count"] == 0
    assert result_data["score_percentage"] == 20.0
    assert result_data["time_taken_seconds"] == 125

    # -------------------------------------------------------------------------
    # STEP 8: Open Result and Review Answers & Explanations
    # -------------------------------------------------------------------------
    review_res = client.get(f"/api/attempts/{attempt_id}/result")
    assert review_res.status_code == 200
    review_data = review_res.json()
    assert len(review_data["question_reviews"]) == 5
    for review in review_data["question_reviews"]:
        assert "correct_index" in review
        assert "selected_index" in review
        assert "explanation" in review
        assert review["explanation"] is not None

    # Verify duplicate submission is rejected
    dup_res = client.post(
        f"/api/attempts/{attempt_id}/submit",
        json={"responses": responses_payload, "time_taken_seconds": 100}
    )
    assert dup_res.status_code == 400
    assert "already been submitted" in dup_res.json()["detail"].lower()

    # -------------------------------------------------------------------------
    # STEP 9: Open Attempt History
    # -------------------------------------------------------------------------
    history_res = client.get("/api/attempts")
    assert history_res.status_code == 200
    history_items = history_res.json()
    assert isinstance(history_items, list)
    matching_attempt = [a for a in history_items if a["id"] == attempt_id]
    assert len(matching_attempt) == 1
    assert matching_attempt[0]["score_percentage"] == 20.0

    # -------------------------------------------------------------------------
    # STEP 10: Open Competency Analysis (Verify Real Quiz Results)
    # -------------------------------------------------------------------------
    comp_res = client.get("/api/competencies?user_id=student_default")
    assert comp_res.status_code == 200
    comp_data = comp_res.json()
    assert comp_data["has_data"] is True
    assert comp_data["total_quizzes_taken"] >= 1
    # Because 1/5 = 20% (< 50% threshold with 5 questions >= 3),
    # this topic should be flagged as a competency gap:
    gap_topics = [g["topic"] for g in comp_data["gaps"]]
    assert "Virtualization & Containers" in gap_topics

    # -------------------------------------------------------------------------
    # STEP 11: Open Personalized Recommendations (Verify Based on Gaps)
    # -------------------------------------------------------------------------
    rec_res = client.get("/api/recommendations?user_id=student_default")
    assert rec_res.status_code == 200
    rec_data = rec_res.json()
    assert rec_data["has_gaps"] is True
    assert len(rec_data["recommendations"]) > 0

    # Check that recommendation targets the weak topic and matches the uploaded material
    matching_rec = [r for r in rec_data["recommendations"] if r["topic"] == "Virtualization & Containers"]
    assert len(matching_rec) >= 1
    assert matching_rec[0]["priority"] == "high"  # 20% accuracy is < 35%, so HIGH priority
    assert matching_rec[0]["has_matching_material"] is True
    assert matching_rec[0]["matched_material_title"] == "Distributed Cloud Systems"

    # -------------------------------------------------------------------------
    # STEP 12: Backend Restart / Persistence Verification
    # -------------------------------------------------------------------------
    # Test persistence by closing the current session, re-querying directly with raw SQL
    with SessionLocal() as db:
        mat_count = db.execute(text("SELECT COUNT(*) FROM learning_materials WHERE id = :id"), {"id": material_id}).scalar()
        assert mat_count == 1, "Learning material did not persist!"

        quiz_count = db.execute(text("SELECT COUNT(*) FROM quizzes WHERE id = :id"), {"id": quiz_id}).scalar()
        assert quiz_count == 1, "Quiz record did not persist!"

        att_count = db.execute(text("SELECT COUNT(*) FROM quiz_attempts WHERE id = :id"), {"id": attempt_id}).scalar()
        assert att_count == 1, "Quiz attempt record did not persist!"

        resp_count = db.execute(text("SELECT COUNT(*) FROM learner_responses WHERE attempt_id = :id"), {"id": attempt_id}).scalar()
        assert resp_count == 5, "Learner responses did not persist!"
