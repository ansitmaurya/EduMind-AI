import io
import re
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Tuple
from backend.app.config import settings

# Attempt import of PyMuPDF (fitz) or fallback to pypdf
try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    fitz = None
    HAS_PYMUPDF = False

try:
    import pypdf
    HAS_PYPDF = True
except ImportError:
    pypdf = None
    HAS_PYPDF = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    docx = None
    HAS_DOCX = False


@dataclass
class ProcessingResult:
    storage_path: str
    file_type: str
    file_size_bytes: int
    extracted_text: str
    extracted_char_count: int
    status: str  # 'ready', 'empty_text', 'failed'
    error_message: Optional[str] = None


class DocumentProcessingError(Exception):
    """Raised when document validation or extraction fails."""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def validate_file_metadata(filename: str, file_size: int) -> str:
    """
    Validates file extension and size constraints.
    Returns the normalized extension.
    """
    if not filename or "." not in filename:
        raise DocumentProcessingError(
            "Invalid file: Filename must include a valid extension (.pdf, .docx, .txt).",
            status_code=400
        )

    ext = filename.rsplit(".", 1)[-1].lower().strip()
    if ext not in settings.ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(settings.ALLOWED_EXTENSIONS))
        raise DocumentProcessingError(
            f"Unsupported file type '.{ext}'. Supported formats: {allowed}.",
            status_code=400
        )

    if file_size <= 0:
        raise DocumentProcessingError(
            "Uploaded file is empty (0 bytes). Please upload a valid document.",
            status_code=400
        )

    if file_size > settings.MAX_UPLOAD_SIZE_BYTES:
        max_mb = settings.MAX_UPLOAD_SIZE_BYTES / (1024 * 1024)
        actual_mb = file_size / (1024 * 1024)
        raise DocumentProcessingError(
            f"File size ({actual_mb:.1f} MB) exceeds maximum allowed limit of {max_mb:.0f} MB.",
            status_code=413
        )

    return ext


def validate_file_content_signature(content_bytes: bytes, ext: str) -> None:
    """
    Validates binary header/magic signature to prevent extension-spoofing.
    Never executes file content.
    """
    if ext == "pdf":
        # PDFs must have %PDF- header within first 1024 bytes
        if b"%PDF-" not in content_bytes[:1024]:
            raise DocumentProcessingError(
                "Invalid PDF document: Header signature '%PDF-' missing or corrupted.",
                status_code=400
            )

    elif ext == "docx":
        # DOCX is a zip archive with 'PK\x03\x04' and contains word/document.xml
        if not content_bytes.startswith(b"PK"):
            raise DocumentProcessingError(
                "Invalid DOCX file: Archive signature missing or corrupted.",
                status_code=400
            )
        try:
            with zipfile.ZipFile(io.BytesIO(content_bytes)) as zf:
                filenames = zf.namelist()
                if "word/document.xml" not in filenames:
                    raise DocumentProcessingError(
                        "Invalid DOCX file: 'word/document.xml' not found in document package.",
                        status_code=400
                    )
        except zipfile.BadZipFile:
            raise DocumentProcessingError(
                "Corrupted DOCX document: Failed to parse as valid OpenXML archive.",
                status_code=400
            )

    elif ext == "txt":
        # Check that it is decodable text and does not contain excessive binary null bytes
        try:
            sample = content_bytes[:4096].decode("utf-8")
        except UnicodeDecodeError:
            try:
                sample = content_bytes[:4096].decode("latin-1")
            except Exception:
                raise DocumentProcessingError(
                    "Invalid text document: Content is not valid text encoded.",
                    status_code=400
                )
        if "\x00" in sample:
            raise DocumentProcessingError(
                "Invalid text document: Binary null characters detected.",
                status_code=400
            )


def extract_text_from_pdf(content_bytes: bytes, file_path: Path) -> Tuple[str, str, Optional[str]]:
    """
    Extracts text from PDF bytes.
    Returns (extracted_text, status, error_message).
    Uses PyMuPDF if available, otherwise pypdf.
    """
    extracted_text = ""
    error_message = None

    if HAS_PYMUPDF:
        try:
            doc = fitz.open(stream=content_bytes, filetype="pdf")
            pages_text = []
            for page in doc:
                text = page.get_text()
                if text:
                    pages_text.append(text)
            doc.close()
            extracted_text = "\n\n".join(pages_text).strip()
        except Exception as exc:
            error_message = f"PyMuPDF extraction failed: {str(exc)}"

    if not extracted_text and HAS_PYPDF and not error_message:
        try:
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            extracted_text = "\n\n".join(pages_text).strip()
        except Exception as exc:
            error_message = f"PDF extraction failed: {str(exc)}"

    if not HAS_PYMUPDF and not HAS_PYPDF:
        error_message = "No PDF extraction engine installed."

    if error_message:
        return "", "failed", error_message

    if not extracted_text:
        return "", "empty_text", "PDF contains no extractable text layer (e.g. scanned image or image-only pages)."

    return extracted_text, "ready", None


def extract_text_from_docx(content_bytes: bytes) -> Tuple[str, str, Optional[str]]:
    """
    Extracts text from DOCX bytes using python-docx.
    Returns (extracted_text, status, error_message).
    """
    if not HAS_DOCX:
        return "", "failed", "python-docx library is not available on system."

    try:
        doc = docx.Document(io.BytesIO(content_bytes))
        paragraphs_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs_text.append(p.text.strip())

        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs_text.append(" | ".join(row_text))

        extracted_text = "\n\n".join(paragraphs_text).strip()

        if not extracted_text:
            return "", "empty_text", "DOCX contains no text in paragraphs or tables."

        return extracted_text, "ready", None
    except Exception as exc:
        return "", "failed", f"DOCX extraction error: {str(exc)}"


def extract_text_from_txt(content_bytes: bytes) -> Tuple[str, str, Optional[str]]:
    """
    Extracts text from TXT bytes with multi-encoding fallback.
    Returns (extracted_text, status, error_message).
    """
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = content_bytes.decode(encoding).strip()
            if not text:
                return "", "empty_text", "Text document is empty."
            return text, "ready", None
        except UnicodeDecodeError:
            continue
    return "", "failed", "Unable to decode text file using supported character encodings."


def clean_extracted_text(text: str) -> str:
    """Normalizes excessive blank lines and whitespace."""
    if not text:
        return ""
    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse 3+ consecutive newlines to 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def process_uploaded_document(filename: str, content_bytes: bytes) -> ProcessingResult:
    """
    Full pipeline:
    1. Validates extension and file size.
    2. Validates binary content signature.
    3. Saves file with safe generated UUID filename.
    4. Extracts text according to file format.
    5. Returns structured ProcessingResult.
    """
    file_size = len(content_bytes)
    ext = validate_file_metadata(filename, file_size)
    validate_file_content_signature(content_bytes, ext)

    # Generate safe unique filename
    safe_filename = f"{uuid.uuid4().hex}.{ext}"
    dest_path = settings.UPLOAD_DIR / safe_filename

    # Save original document to disk safely
    with open(dest_path, "wb") as f:
        f.write(content_bytes)

    # Relative path stored in database for portability
    rel_storage_path = str(dest_path.relative_to(settings.BASE_DIR))

    # Extract text based on file type
    if ext == "pdf":
        raw_text, status, error_message = extract_text_from_pdf(content_bytes, dest_path)
    elif ext == "docx":
        raw_text, status, error_message = extract_text_from_docx(content_bytes)
    elif ext == "txt":
        raw_text, status, error_message = extract_text_from_txt(content_bytes)
    else:
        raw_text, status, error_message = "", "failed", f"Unsupported extension: {ext}"

    cleaned_text = clean_extracted_text(raw_text)
    char_count = len(cleaned_text)

    return ProcessingResult(
        storage_path=rel_storage_path,
        file_type=ext,
        file_size_bytes=file_size,
        extracted_text=cleaned_text,
        extracted_char_count=char_count,
        status=status,
        error_message=error_message
    )
